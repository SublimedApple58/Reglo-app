"""
Round pre-seed — esporta i file dell'outreach (fuori dal repo) in un JSON grezzo
che `scripts/seed-round-preseed.ts` legge. Nessuna logica di prodotto qui: solo
celle. I dati sono riservati: il JSON va nello scratchpad, MAI nel repo.

  python3 -I scripts/round-preseed/export-outreach.py \
    ~/Developer/reglo-deck-raccolta-2026/outreach  /percorso/outreach.json
"""
import csv
import json
import re
import sys
from pathlib import Path

import openpyxl

src = Path(sys.argv[1]).expanduser()
out = Path(sys.argv[2]).expanduser()


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


contacts = []
wb = openpyxl.load_workbook(src / "contatti-investitori.xlsx", data_only=True)
for ws in wb.worksheets[1:]:
    group = None
    for row in ws.iter_rows(values_only=True):
        first = clean(row[0])
        if first and re.search(r"\(\d+\)$", first) and not clean(row[1]):
            group = re.sub(r"\s*\(\d+\)$", "", first)
            continue
        if not first or first == "Nome":
            continue
        contacts.append(
            {
                "sheet": ws.title,
                "group": group,
                "name": first,
                "organization": clean(row[1]),
                "channel": clean(row[2]),
                "status": clean(row[3]),
                "notes": clean(row[4]),
            }
        )

with open(src / "archivio" / "investitori.csv", newline="") as fh:
    archive = [{k: clean(v) for k, v in r.items()} for r in csv.DictReader(fh)]

angels = []
wa = openpyxl.load_workbook(src / "archivio" / "angel-ondata3.xlsx", data_only=True)
first_sheet = wa.worksheets[0]
header = [clean(c) for c in next(first_sheet.iter_rows(values_only=True))]
for row in first_sheet.iter_rows(min_row=2, values_only=True):
    item = {header[i]: clean(v) for i, v in enumerate(row) if i < len(header) and header[i]}
    if item.get("Nome"):
        angels.append(item)

def sheet_pairs(name):
    if name not in wa.sheetnames:
        return []
    rows = list(wa[name].iter_rows(min_row=2, values_only=True))
    return [{"name": clean(r[0]), "reason": clean(r[1])} for r in rows if clean(r[0])]


angel_conflicts = sheet_pairs("Esclusi per conflitto")
angel_paused = sheet_pairs("In pausa o da verificare")

leads = []
wl = openpyxl.load_workbook(src / "fondi-lead-italia.xlsx", data_only=True)
ls = wl.worksheets[0]
lh = [clean(c) for c in next(ls.iter_rows(values_only=True))]
for row in ls.iter_rows(min_row=2, values_only=True):
    item = {lh[i]: clean(v) for i, v in enumerate(row) if i < len(lh) and lh[i]}
    if item.get("Fondo"):
        leads.append(item)

out.write_text(json.dumps({"contacts": contacts, "archive": archive, "angels": angels, "angelConflicts": angel_conflicts, "angelPaused": angel_paused, "leads": leads}, ensure_ascii=False, indent=1))
print(f"{len(contacts)} contatti, {len(archive)} schede archivio, {len(angels)} angel, {len(leads)} fondi lead → {out}")
